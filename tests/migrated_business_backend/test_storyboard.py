"""Synthetic real DOCX and direct aiohttp handler checks; no listening server/model."""
import asyncio
import copy
import hashlib
import io
import json
from pathlib import Path
import socket
import sys
import tempfile
import types
import unittest
from unittest.mock import patch
import zipfile

from aiohttp import web
from docx import Document
from docx.shared import Inches
from PIL import Image

from backend.assets import AssetCatalog
from backend.storyboard.document_import import parse_storyboard_document, preview_storyboard_document
from backend.storyboard.script_document import parse_docx, normalize_inventory
from backend.storyboard.script_contract import storyboard_projection
from backend.storyboard.services import StoryboardServices, ProviderUnavailable
from backend.storyboard.table_adapter import validated_table
from backend.storyboard.table_source import project_selected_source
from backend.prompt.local_media import LocalMediaResolver
from backend.prompt.prompt_service import parse_request, prepare_context
from backend.prompt.prompt_compiler import compile_prompt
from backend.integrations.comfyui.business_nodes import register_routes


def docx_fixture():
    """Two logical tasks, repeated identical image occurrences and a vertical merge."""
    image = io.BytesIO()
    Image.new('RGB', (12, 8), '#004477').save(image, format='PNG')
    document = Document()
    document.add_paragraph('第一篇章 合成测试，无真实用户资料')
    table = document.add_table(rows=3, cols=4)
    for cell, text in zip(table.rows[0].cells, ('镜号', '画面', '参考图', '备注')):
        cell.text = text
    for n in (1, 2):
        table.cell(n, 0).text = str(n)
        table.cell(n, 1).text = f'画面 {n}：几何形状。\n保持原文。'
        table.cell(n, 2).paragraphs[0].add_run().add_picture(io.BytesIO(image.getvalue()), width=Inches(.2))
    table.cell(1, 3).text = '连续镜头备注'
    table.cell(1, 3).merge(table.cell(2, 3))
    title = document.add_table(rows=1, cols=2)
    title.cell(0, 0).merge(title.cell(0, 1)).text = '第二篇章（横向合并）'
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def choices_for(inventory):
    return {'tables': {t['id']: {'mapping': t['mapping'], 'rows': {}}
                       for t in inventory['tables'] if t['header']}, 'images': {}}


def prompt_table():
    return {'fields': [{'id': 'scene', 'name': '画面', 'type': 'text'},
                       {'id': 'ref', 'name': '参考图', 'type': 'assets', 'video_reference': True}],
            'records': [{'id': 'row-1', 'values': {'scene': '一个蓝色立方体', 'ref': [
                {'id': 'asset-1', 'url': '/view?filename=fixture.png&type=input'}]}}],
            'meta': {'video_reference_version': 1, 'prompt_config': {'bindings': {'image_prompt': 'scene', 'image_url': 'ref'}}}}


class Request:
    def __init__(self, payload=None, *, file=None, query=None):
        self.payload = json.dumps(payload).encode() if not isinstance(payload, bytes) else payload
        self.content = self
        self.query = query or {}
        self.file = file
    async def iter_chunked(self, size):
        for start in range(0, len(self.payload), size):
            yield self.payload[start:start + size]
    async def multipart(self):
        return self
    async def next(self):
        if self.file is None:
            return None
        filename, data = self.file
        self.file = None
        class Part:
            name = 'file'
            def __init__(self):
                self.filename = filename
                self.stream = io.BytesIO(data)
            async def read_chunk(self, size=65536):
                return self.stream.read(size)
        return Part()


class BusinessTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.docx = docx_fixture()

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.input = self.root / 'input'
        self.input.mkdir()
        Image.new('RGB', (2, 2), '#ff9933').save(self.input / 'fixture.png')
        self.services = StoryboardServices(cache_directory=self.root / 'cache', input_directory=self.input,
                                          local_media=LocalMediaResolver(roots={'input': self.input}))

    def test_docx_merges_occurrences_and_scene_are_lossless(self):
        inventory, blobs = parse_docx(self.docx, '合成.docx')
        first = inventory['tables'][0]
        self.assertEqual(first['mapping']['1'], 'scene')
        anchor = first['rows'][1]['cells'][-1]
        self.assertEqual(anchor['rowspan'], 2)
        self.assertEqual(first['rows'][2]['cells'][-1]['anchor'], anchor['id'])
        self.assertEqual(inventory['tables'][1]['rows'][0]['cells'][0]['colspan'], 2)
        self.assertEqual(len(blobs), 1)
        self.assertEqual(len(inventory['occurrences']), 2)
        result = normalize_inventory(inventory, choices_for(inventory))
        self.assertEqual(len(result['tasks']), 2)
        self.assertEqual(result['tasks'][1]['values']['notes'], '连续镜头备注')
        self.assertEqual(result['tasks'][0]['values']['scene'], '画面 1：几何形状。\n保持原文。')
        self.assertNotEqual(result['tasks'][0]['images'][0]['id'], result['tasks'][1]['images'][0]['id'])

    def test_preview_compatible_dto_and_visual_column_not_reference(self):
        result = preview_storyboard_document('合成.docx', self.docx)
        self.assertEqual(result['tables'][0]['mapping']['image_prompt'], 1)
        self.assertEqual(result['tables'][0]['mapping']['reference_image'], 2)
        self.assertEqual(len(result['assets']), 1)
        self.assertIn('1:2', result['tables'][0]['images'])
        legacy = parse_storyboard_document('合成.docx', self.docx)
        self.assertEqual(len(legacy['shots']), 2)
        self.assertTrue(legacy['shots'][0]['prompt'].startswith('画面 1'))

    def test_repeated_import_preserves_original_and_deduplicates_assets(self):
        catalog = AssetCatalog(self.root / 'db.sqlite', roots={'input': self.input})
        self.services.asset_catalog = catalog
        first = self.services.import_document(self.docx, '合成.docx')
        second = self.services.import_document(self.docx, '合成.docx')
        self.assertEqual(first, second)
        self.assertEqual(len(list((self.input / 'DAELAB/script-parser').glob('*.png'))), 1)
        self.assertEqual((self.root / 'cache' / (first['document_id'] + '.docx')).read_bytes(), self.docx)
        stored = catalog.get(next(iter(first['assets'].values()))['catalog_asset_id'])
        self.assertEqual(len(stored['references']), 1)
        self.assertEqual(len(stored['provenance']), 1)
        with catalog.database.read() as connection:
            self.assertEqual(connection.execute('SELECT COUNT(*) FROM assets').fetchone()[0], 1)
        resolver = LocalMediaResolver(catalog=catalog)
        media = resolver(next(iter(first['assets'].values()))['url'])
        self.assertEqual(media['kind'], 'image')
        recreated = StoryboardServices(cache_directory=self.root / 'cache', input_directory=self.input)
        self.assertEqual(recreated.read_inventory(first['document_id']), first)

    def test_normalize_accepts_saved_inventory_and_rejects_identity_mismatch(self):
        inventory = self.services.import_document(self.docx, '合成.docx')
        body = {'document_id': inventory['document_id'], 'inventory': inventory, 'choices': choices_for(inventory)}
        self.assertEqual(len(self.services.normalize(body)['tasks']), 2)
        body['document_id'] = '0' * 64
        with self.assertRaisesRegex(ValueError, '不匹配'):
            self.services.normalize(body)

    def test_oversize_table_is_rejected_without_slicing(self):
        for table in ({'fields': [{'id': str(n)} for n in range(65)], 'records': []},
                      {'fields': [], 'records': [{'id': str(n), 'values': {}} for n in range(501)]}):
            with self.assertRaisesRegex(ValueError, '未截断'):
                validated_table(table)
        csv = ('镜号,画面\n' + ''.join(f'{n},画面{n}\n' for n in range(501))).encode()
        with self.assertRaisesRegex(ValueError, '未截断'):
            preview_storyboard_document('overflow.csv', csv)
        numbered = '\n'.join(f'{n}. 画面' for n in range(1, 502)).encode()
        with self.assertRaisesRegex(ValueError, '未截断'):
            preview_storyboard_document('overflow.txt', numbered)

    def test_exact_limit_and_unknown_table_metadata_survive(self):
        value = {'fields': [{'id': f'f{n}'} for n in range(64)],
                 'records': [{'id': f'r{n}', 'values': {}, 'future': {'raw': True}} for n in range(500)],
                 'meta': {'future': [1, 2]}, 'extension': {'preserve': 'yes'}}
        self.assertEqual(validated_table(value), value)
        self.assertIsNot(validated_table(value), value)

    def test_archive_traversal_duplicate_and_entities_are_rejected(self):
        for name, data in (('../bad.xml', b''), ('word/document.xml', b'<bad/>')):
            stream = io.BytesIO(self.docx)
            with zipfile.ZipFile(stream, 'a') as archive:
                archive.writestr(name, data)
            with self.assertRaises(ValueError):
                parse_docx(stream.getvalue())
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, 'w') as archive:
            archive.writestr('word/document.xml', '<!DOCTYPE x [<!ENTITY y "danger">]><x/>')
        with self.assertRaisesRegex(ValueError, '外部实体'):
            parse_docx(stream.getvalue())

    def test_paths_block_traversal_absolute_cloud_and_duplicate_query(self):
        for url in ('https://example.invalid/a.png', '/view?filename=../fixture.png',
                    '/view?filename=C:%5Csecret.png', '/view?filename=a.png&subfolder=../',
                    '/view?filename=a.png&filename=fixture.png', '/view?filename=fixture.png&type=user'):
            with self.assertRaises(ValueError):
                self.services.resolve_media(url)
        with self.assertRaises(ValueError):
            self.services.read_inventory('../../outside')

    def test_prompt_parse_validate_review_and_stale_file(self):
        table = prompt_table()
        payload = {'table': table, 'recordIds': ['row-1']}
        result = parse_request(payload, self.services.resolve_media)['results'][0]
        self.assertNotIn('path', result['assets'][0])
        self.assertEqual(result['assets'][0]['sha256'], hashlib.sha256((self.input / 'fixture.png').read_bytes()).hexdigest())
        payload.update(operation='validate', documents={'row-1': result['document']})
        self.assertNotIn('error', parse_request(payload, self.services.resolve_media)['results'][0])
        Image.new('RGB', (2, 2), 'blue').save(self.input / 'fixture.png')
        self.assertIn('error', parse_request(payload, self.services.resolve_media)['results'][0])
        payload['operation'] = 'review'
        reviewed = parse_request(payload, self.services.resolve_media)['results'][0]
        self.assertNotIn('error', reviewed)
        self.assertNotEqual(reviewed['document']['sourceFingerprint'], result['document']['sourceFingerprint'])

    def test_wrong_visual_type_cannot_be_used_as_prompt(self):
        table = prompt_table()
        table['fields'][0]['type'] = 'assets'
        result = parse_request({'table': table, 'recordIds': ['row-1']}, self.services.resolve_media)
        self.assertIn('类型已变化', result['results'][0]['error'])

    def test_stable_column_template_and_reference_order(self):
        table = prompt_table()
        template = {'kind': 'column-template', 'version': 1, 'segments': [
            {'type': 'column', 'fieldId': 'scene', 'valueType': 'text'},
            {'type': 'column', 'fieldId': 'ref', 'assetId': 'asset-1'}]}
        table['fields'].append({'id': 'prompt', 'name': '提示词', 'type': 'json', 'presentation': 'prompt', 'promptTemplate': template})
        table['meta']['prompt_config']['bindings']['final_prompt'] = 'prompt'
        result = parse_request({'table': table, 'recordIds': ['row-1']}, self.services.resolve_media)['results'][0]
        context = prepare_context(table, table['records'][0], self.services.resolve_media)
        text, assets = compile_prompt(result['document'], context)
        self.assertEqual(text, '一个蓝色立方体@image_1')
        self.assertEqual(assets[0]['assetId'], 'asset-1')
        self.assertEqual(len(text.reference_spans), 1)

    def test_selected_source_rows_follow_identity_not_index(self):
        source = {'fields': [{'id': 'a', 'type': 'text'}], 'records': [
            {'id': 'r2', 'values': {'a': 'second'}}, {'id': 'r1', 'values': {'a': 'first'}}]}
        local = {'fields': [{'id': 'b'}], 'records': [{'id': 'r1', 'selected': True, 'values': {}}],
                 'meta': {'liveSource': {'fields': {'b': 'a'}}}}
        result = project_selected_source(local, source)
        self.assertEqual(result['records'][0]['id'], 'r1')
        self.assertEqual(result['records'][0]['values'], {'b': 'first'})

    def test_routes_are_explicit_atomic_idempotent_and_offline(self):
        server = types.SimpleNamespace(routes=web.RouteTableDef())
        self.assertEqual(len(server.routes), 0)
        self.assertEqual(len(register_routes(server, self.services)), 6)
        self.assertEqual(register_routes(server, self.services), ())
        with self.assertRaisesRegex(RuntimeError, 'duplicate'):
            register_routes(server, copy.copy(self.services))
        other = types.SimpleNamespace(routes=web.RouteTableDef())
        other.routes.post('/daelab/script-parser/classify')(lambda request: None)
        with self.assertRaises(RuntimeError):
            register_routes(other, self.services)
        self.assertEqual(len(other.routes), 1)
        handlers = {r.path: r.handler for r in server.routes}
        loop = asyncio.new_event_loop()
        try:
            # Blocking remote connection methods preserves Windows asyncio's local wakeup sockets.
            with patch.object(socket.socket, 'connect', side_effect=AssertionError('network forbidden')), \
                 patch.object(socket.socket, 'connect_ex', side_effect=AssertionError('network forbidden')):
                response = loop.run_until_complete(handlers['/daelab/script-parser/classify'](Request({})))
                self.assertEqual(response.status, 503)
                status = loop.run_until_complete(handlers['/daelab/script-parser/ai-status'](Request({})))
                self.assertFalse(json.loads(status.body)['configured'])
        finally:
            loop.close()

    def test_routes_roundtrip_upload_preview_normalize_and_parse(self):
        server = types.SimpleNamespace(routes=web.RouteTableDef())
        register_routes(server, self.services)
        handlers = {r.path: r.handler for r in server.routes}
        async def exercise():
            response = await handlers['/daelab/script-parser/document'](Request(file=('fixture.docx', self.docx)))
            self.assertEqual(response.status, 200)
            inventory = json.loads(response.body)
            response = await handlers['/daelab/script-parser/normalize'](Request({'document_id': inventory['document_id'], 'choices': choices_for(inventory)}))
            self.assertEqual(len(json.loads(response.body)['tasks']), 2)
            response = await handlers['/daelab/storyboard/import_document'](Request(file=('fixture.docx', self.docx), query={'preview': '1'}))
            self.assertEqual(len(json.loads(response.body)['assets']), 1)
            response = await handlers['/daelab/storyboard/parse-prompts'](Request({'table': prompt_table(), 'recordIds': ['row-1']}))
            self.assertEqual(response.status, 200)
            self.assertNotIn('error', json.loads(response.body)['results'][0])
        asyncio.run(exercise())

    def test_classifier_fixture_must_cover_every_cell_and_cannot_rewrite(self):
        inventory = self.services.import_document(self.docx, 'fixture.docx')
        body = {'document_id': inventory['document_id'], 'choices': choices_for(inventory)}
        with self.assertRaises(ProviderUnavailable):
            asyncio.run(self.services.classify(body))
        def fixture(rows):
            return {'rows': [{'row_id': r['row_id'], 'kind': 'review',
                             'cells': [{'cell_id': c['cell_id'], 'role': 'other'} for c in r['cells']]} for r in rows]}
        self.services.classifier = fixture
        result = asyncio.run(self.services.classify(body))
        self.assertEqual(len(result['tables']['t1']['rows']), 3)
        self.services.classifier = lambda rows: {'rows': []}
        with self.assertRaisesRegex(ValueError, '遗漏'):
            asyncio.run(self.services.classify(body))
        def bad(rows):
            value = fixture(rows)
            value['rows'][0]['rewrite'] = 'invented'
            return value
        self.services.classifier = bad
        with self.assertRaisesRegex(ValueError, '改写'):
            asyncio.run(self.services.classify(body))

    def test_provider_failure_is_sanitized_and_no_success_is_fabricated(self):
        inventory = self.services.import_document(self.docx, 'fixture.docx')
        def broken(rows):
            raise ValueError('secret-from-provider-response')
        self.services.classifier = broken
        server = types.SimpleNamespace(routes=web.RouteTableDef())
        register_routes(server, self.services)
        handler = next(r.handler for r in server.routes if r.path.endswith('/classify'))
        response = asyncio.run(handler(Request({'document_id': inventory['document_id'], 'choices': choices_for(inventory)})))
        self.assertEqual(response.status, 502)
        self.assertNotIn('secret-from-provider-response', response.text)
        self.assertNotIn('choices', json.loads(response.body))

    def test_missing_media_port_is_an_explicit_unavailable_response(self):
        self.services.local_media = None
        server = types.SimpleNamespace(routes=web.RouteTableDef())
        register_routes(server, self.services)
        handler = next(r.handler for r in server.routes if r.path.endswith('/parse-prompts'))
        response = asyncio.run(handler(Request({'table': prompt_table(), 'recordIds': ['row-1']})))
        self.assertEqual(response.status, 503)
        self.assertEqual(json.loads(response.body)['code'], 'service_unavailable')

    def test_nodes_exact_identity_and_native_slot_order(self):
        from backend.integrations.comfyui.business_nodes import register_nodes
        class String:
            @staticmethod
            def Input(name, **kwargs): return types.SimpleNamespace(name=name, **kwargs)
            @staticmethod
            def Output(name, **kwargs): return types.SimpleNamespace(name=name, **kwargs)
        fake = types.SimpleNamespace(ComfyNode=object, String=String,
                                     Schema=lambda **kwargs: types.SimpleNamespace(**kwargs), NodeOutput=lambda *args: args)
        with patch.dict(sys.modules, {'comfy_api': types.ModuleType('comfy_api'), 'comfy_api.latest': types.SimpleNamespace(io=fake)}):
            mappings, names = register_nodes()
            self.assertEqual(set(mappings), {'DAELAB.Table', 'DAELAB.ScriptParser', 'DAELAB.StoryboardImport', 'DAELAB.LLMStylePreviewB'})
            self.assertEqual(set(mappings), set(names))
            schema = mappings['DAELAB.Table'].define_schema()
            self.assertEqual([s.name for s in schema.inputs], ['table_data', 'source_table'])
            schema = mappings['DAELAB.ScriptParser'].define_schema()
            self.assertEqual([s.name for s in schema.outputs], ['table_json', 'storyboard_json'])
            result = mappings['DAELAB.Table'].execute(json.dumps({'fields': [], 'records': [], 'future': {'keep': 1}}))
            self.assertEqual(json.loads(result[0])['future'], {'keep': 1})
            self.assertEqual(mappings['DAELAB.LLMStylePreviewB']().output('approved'), ('approved', ''))
            with self.assertRaises(ValueError):
                mappings['DAELAB.LLMStylePreviewB']().output('')


if __name__ == '__main__':
    unittest.main()
